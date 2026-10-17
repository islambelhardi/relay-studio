from io import BytesIO
import pytest
from docx import Document
from pypdf import PdfReader, PdfWriter
from reportlab.pdfgen.canvas import Canvas
from contentops.extraction import ExtractionError, MAX_BYTES, extract_document

@pytest.mark.parametrize("ext",["txt","md"])
def test_text_formats(ext):
    assert extract_document(b"\xef\xbb\xbf# Product requirements",ext) == "# Product requirements"

def test_docx_includes_tables():
    doc = Document()
    doc.add_paragraph("Planned beta")
    doc.add_table(rows=1,cols=1).cell(0,0).text = "Invite only"
    stream = BytesIO()
    doc.save(stream)
    assert "Invite only" in extract_document(stream.getvalue(),"docx")

def test_text_pdf():
    stream = BytesIO()
    canvas = Canvas(stream)
    canvas.drawString(50,700,"Planned private beta")
    canvas.save()
    assert "Planned private beta" in extract_document(stream.getvalue(),"pdf")

def test_encrypted_and_image_only_pdf():
    stream = BytesIO()
    writer = PdfWriter()
    writer.add_blank_page(width=100,height=100)
    writer.write(stream)
    with pytest.raises(ExtractionError,match="selectable text"):
        extract_document(stream.getvalue(),"pdf")
    encrypted = BytesIO()
    writer.encrypt("secret")
    writer.write(encrypted)
    with pytest.raises(ExtractionError,match="password"):
        extract_document(encrypted.getvalue(),"pdf")

@pytest.mark.parametrize("data,ext",[(b"","md"),(b"  ","txt"),(b"broken","pdf"),(b"broken","docx"),(b"abc","exe"),(b"\xff","txt"),(b"a\x00b","md"),(b"x"*(MAX_BYTES+1),"md")],ids=['empty','whitespace','corrupt-pdf','corrupt-docx','unsupported','bad-encoding','binary-text','oversized'])
def test_invalid_uploads(data,ext):
    with pytest.raises(ExtractionError):
        extract_document(data,ext)
