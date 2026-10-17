from io import BytesIO
from pathlib import Path
from zipfile import ZipFile, BadZipFile
from pypdf import PdfReader
from docx import Document

MAX_BYTES = 10 * 1024 * 1024

class ExtractionError(ValueError):
    pass

def extract_document(data: bytes, extension: str) -> str:
    if not data or len(data) > MAX_BYTES:
        raise ExtractionError("Choose a nonempty document smaller than 10 MB.")
    extension = extension.lower().lstrip(".")
    try:
        if extension in {"txt", "md"}:
            text = data.decode("utf-8-sig")
            if "\x00" in text:
                raise ExtractionError("The file is not readable text. Export it as UTF-8 TXT or Markdown.")
        elif extension == "pdf":
            if not data.startswith(b"%PDF-"):
                raise ExtractionError("This PDF is corrupt. Export a fresh copy and try again.")
            reader = PdfReader(BytesIO(data), strict=False)
            if reader.is_encrypted:
                raise ExtractionError("Remove the PDF password and upload an unencrypted copy.")
            if len(reader.pages) > 300:
                raise ExtractionError("Use a PDF with at most 300 pages.")
            text = "\n\n".join(page.extract_text() or "" for page in reader.pages)
            if not text.strip():
                raise ExtractionError("This PDF has no selectable text. Paste the PRD or upload a text-based PDF; OCR is not included.")
        elif extension == "docx":
            with ZipFile(BytesIO(data)) as archive:
                if sum(i.file_size for i in archive.infolist()) > 50 * 1024 * 1024:
                    raise ExtractionError("The Word document expands beyond the supported size. Export a smaller copy.")
            document = Document(BytesIO(data))
            text = "\n\n".join([p.text for p in document.paragraphs] + [" | ".join(c.text for c in row.cells) for table in document.tables for row in table.rows])
        else:
            raise ExtractionError("Supported formats: TXT, Markdown, PDF, and DOCX.")
    except ExtractionError:
        raise
    except (UnicodeDecodeError, BadZipFile):
        raise ExtractionError("The document is corrupt or is not in the selected format. Export a fresh copy.") from None
    except Exception:
        raise ExtractionError("The document could not be read. Export a fresh, unencrypted copy or paste its text.") from None
    text = text.strip()
    if not text:
        raise ExtractionError("The document is empty. Add product requirements before starting a run.")
    if len(text) > 200_000:
        raise ExtractionError("The extracted PRD is too long. Keep it below 200,000 characters.")
    return text
